"""DOCX export generator for ATS-compliant Word documents."""

from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor


class CVDocxRenderer:
    """Renders structured Markdown CV text into an ATS-compliant Word (.docx) document."""

    def render_markdown_to_docx_bytes(self, markdown_text: str) -> bytes:
        """Parse markdown lines and render directly to in-memory DOCX bytes."""
        doc = Document()

        # Set 0.5 inch margins for standard ATS format
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

        lines = markdown_text.splitlines()

        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue

            # Level 1: Candidate Name
            if stripped.startswith("# "):
                name = stripped[2:].strip()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(name)
                run.bold = True
                run.font.size = Pt(18)
                run.font.name = "Arial"
                run.font.color.rgb = RGBColor(26, 32, 44)
                continue

            # Sub-title / Role (**Role**)
            if (
                stripped.startswith("**")
                and stripped.endswith("**")
                and not stripped.startswith("**Key")
                and not stripped.startswith("**Technical")
            ):
                title_text = stripped[2:-2].strip()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(title_text)
                run.bold = True
                run.font.size = Pt(11)
                run.font.name = "Arial"
                run.font.color.rgb = RGBColor(49, 130, 206)
                continue

            # Level 2: Section Header (## Section)
            if stripped.startswith("## "):
                section_title = stripped[3:].strip().upper()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(section_title)
                run.bold = True
                run.font.size = Pt(11)
                run.font.name = "Arial"
                run.font.color.rgb = RGBColor(26, 32, 44)
                continue

            # Level 3: Item Header (### Role | Company | Dates)
            if stripped.startswith("### "):
                role_header = stripped[4:].strip()
                clean_header = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1", role_header)
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(1)
                run = p.add_run(clean_header)
                run.bold = True
                run.font.size = Pt(10)
                run.font.name = "Arial"
                run.font.color.rgb = RGBColor(45, 55, 72)
                continue

            # Bullet points
            if stripped.startswith(("- ", "* ", "● ", "▪ ")):
                bullet_content = stripped[2:].strip()
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                self._add_inline_formatted_runs(p, bullet_content)
                continue

            # Contact row or general paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            self._add_inline_formatted_runs(p, stripped)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def _add_inline_formatted_runs(self, paragraph, text: str) -> None:
        """Parse markdown bold (**bold**) and hyperlinks [text](url) into formatted runs."""
        # Simple tokenization for bold markers
        parts = re.split(r"(\*\*.*?\*\*)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                clean_bold = part[2:-2]
                clean_bold = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1", clean_bold)
                run = paragraph.add_run(clean_bold)
                run.bold = True
                run.font.size = Pt(9.5)
                run.font.name = "Arial"
            else:
                clean_text = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1", part)
                run = paragraph.add_run(clean_text)
                run.font.size = Pt(9.5)
                run.font.name = "Arial"

    def render_markdown_to_file(self, markdown_text: str, output_path: Path) -> Path:
        """Render markdown to a DOCX file on disk."""
        docx_bytes = self.render_markdown_to_docx_bytes(markdown_text)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(docx_bytes)
        return output_path
